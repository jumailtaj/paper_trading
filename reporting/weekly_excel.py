"""Weekly Excel report: the persistent trade record for V1 (no database).

reports/weekly_trades.xlsx has one sheet per weekday (Monday..Friday). A new trading week archives
the previous workbook to archive/weekly_trades_YYYY-MM-DD.xlsx (never deleted) and starts a fresh one.
"""
import logging
import os
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import config

log = logging.getLogger("paper_trader")

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
HEADERS = ["Trade ID", "Strategy ID", "Symbol", "Side", "Quantity", "Entry Price", "Exit Price",
           "Entry Time", "Exit Time", "Gross P&L", "Brokerage", "STT", "Transaction Charges",
           "SEBI Charges", "GST", "Stamp Duty", "Total Charges", "Net P&L", "Exit Reason"]
SUMMARY_LABELS = ["Total Trades", "Winning Trades", "Losing Trades", "Gross P&L", "Total Charges", "Net P&L"]
COL = {name: i + 1 for i, name in enumerate(HEADERS)}
WEEK_MARKER = "WEEK_START="
TIME_FORMAT = "yyyy-mm-dd hh:mm:ss"


def week_start(day: date) -> date:
    return day - timedelta(days=day.weekday())


class WeeklyExcel:
    def __init__(self, path: Path = config.WEEKLY_REPORT_PATH, archive_dir: Path = config.ARCHIVE_DIR):
        self.path = Path(path)
        self.archive_dir = Path(archive_dir)

    # ------------------------------------------------------------------ workbook lifecycle
    def _create(self, monday: date) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        wb.remove(wb.active)
        for day in DAYS:
            ws = wb.create_sheet(day)
            self._write_header(ws)
        wb.properties.subject = f"{WEEK_MARKER}{monday.isoformat()}"
        self._save(wb)

    @staticmethod
    def _write_header(ws) -> None:
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for i, h in enumerate(HEADERS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) + 2)
        ws.column_dimensions[get_column_letter(COL["Entry Time"])].width = 20
        ws.column_dimensions[get_column_letter(COL["Exit Time"])].width = 20

    def _save(self, wb) -> None:
        tmp = self.path.with_suffix(".tmp.xlsx")
        wb.save(tmp)
        os.replace(tmp, self.path)  # atomic: never leaves a half-written report

    def _file_week(self) -> Optional[date]:
        wb = load_workbook(self.path)
        subject = wb.properties.subject or ""
        if subject.startswith(WEEK_MARKER):
            return date.fromisoformat(subject[len(WEEK_MARKER):])
        return week_start(datetime.fromtimestamp(self.path.stat().st_mtime).date())

    def ensure_week(self, for_date: date) -> None:
        """Make sure the active workbook belongs to the week of `for_date`; archive the old one if not."""
        monday = week_start(for_date)
        if not self.path.exists():
            self._create(monday)
            return
        file_week = self._file_week()
        if file_week == monday:
            return
        self.archive_dir.mkdir(parents=True, exist_ok=True)
        target = self.archive_dir / f"weekly_trades_{file_week.isoformat()}.xlsx"
        n = 1
        while target.exists():  # never overwrite history
            target = self.archive_dir / f"weekly_trades_{file_week.isoformat()}_{n}.xlsx"
            n += 1
        os.replace(self.path, target)
        log.info("WEEKLY REPORT ARCHIVED %s", target.name)
        self._create(monday)

    # ------------------------------------------------------------------ rows
    @staticmethod
    def _read_trade_rows(ws) -> List[list]:
        rows = []
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value in (None, ""):
                break  # blank row = start of the summary block
            rows.append([ws.cell(r, c).value for c in range(1, len(HEADERS) + 1)])
        return rows

    @staticmethod
    def _trade_to_row(trade) -> list:
        return [trade.trade_id, trade.strategy_id, trade.symbol, trade.side, trade.quantity,
                trade.entry_price, trade.exit_price, trade.entry_time, trade.exit_time,
                trade.gross_pnl, trade.brokerage, trade.stt, trade.transaction_charges,
                trade.sebi_charges, trade.gst, trade.stamp_duty, trade.total_charges,
                trade.net_pnl, trade.exit_reason]

    @staticmethod
    def summarize(rows: List[list]) -> dict:
        net = [r[COL["Net P&L"] - 1] or 0 for r in rows]
        return {
            "Total Trades": len(rows),
            "Winning Trades": sum(1 for n in net if n > 0),
            "Losing Trades": sum(1 for n in net if n < 0),
            "Gross P&L": round(sum(r[COL["Gross P&L"] - 1] or 0 for r in rows), 2),
            "Total Charges": round(sum(r[COL["Total Charges"] - 1] or 0 for r in rows), 2),
            "Net P&L": round(sum(net), 2),
        }

    def save_trade(self, trade) -> None:
        """Append a completed trade to its weekday sheet and refresh that sheet's totals."""
        day = trade.exit_time.date()
        if day.weekday() > 4:
            raise ValueError(f"Cannot record a trade on a weekend ({day}); the report has Mon-Fri sheets only")
        self.ensure_week(day)
        wb = load_workbook(self.path)
        ws = wb[DAYS[day.weekday()]]
        rows = self._read_trade_rows(ws)
        rows.append(self._trade_to_row(trade))

        ws.delete_rows(2, ws.max_row)
        for row in rows:
            ws.append(row)
        for r in range(2, len(rows) + 2):
            for name in ("Entry Time", "Exit Time"):
                ws.cell(r, COL[name]).number_format = TIME_FORMAT
        summary_row = len(rows) + 3  # one blank row between trades and totals
        for i, (label, value) in enumerate(self.summarize(rows).items()):
            ws.cell(summary_row + i, 1, label).font = Font(bold=True)
            ws.cell(summary_row + i, 2, value)
        self._save(wb)
        log.info("EXCEL UPDATED sheet=%s trade=%s", DAYS[day.weekday()], trade.trade_id)

    def count_trades(self, day: date) -> int:
        """Trades already recorded for `day` (0 if the report is missing or belongs to another week)."""
        if day.weekday() > 4 or not self.path.exists():
            return 0
        if self._file_week() != week_start(day):
            return 0
        ws = load_workbook(self.path)[DAYS[day.weekday()]]
        return len(self._read_trade_rows(ws))
