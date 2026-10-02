from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from reporting.weekly_excel import DAYS, HEADERS, WeeklyExcel
from trading.charge_calculator import ChargeCalculator
from trading.positions import Trade


def make_trade(exit_time, entry=1000.0, exit_=1020.0, side="BUY", tid="T1", symbol="RELIANCE"):
    gross = (exit_ - entry) * 10 if side == "BUY" else (entry - exit_) * 10
    t = Trade(tid, "EMA-01", symbol, side, 10, entry, exit_, exit_time.replace(hour=9, minute=40),
              exit_time, gross, "TARGET")
    t.apply_charges(ChargeCalculator().calculate(t))
    return t


@pytest.fixture
def xl(tmp_path):
    return WeeklyExcel(tmp_path / "reports" / "weekly_trades.xlsx", tmp_path / "archive")


WED = datetime(2026, 9, 30, 10, 30)   # Wednesday
THU = datetime(2026, 10, 1, 11, 0)    # Thursday


def rows_of(ws):
    return [[c.value for c in r] for r in ws.iter_rows(min_row=1)]


def test_workbook_has_monday_to_friday_sheets_with_headers(xl):
    xl.ensure_week(date(2026, 9, 30))
    wb = load_workbook(xl.path)
    assert wb.sheetnames == DAYS
    assert [c.value for c in wb["Monday"][1]] == HEADERS
    assert "Exit Reason" in HEADERS and "Net P&L" in HEADERS


def test_trade_written_to_correct_weekday_with_all_columns(xl):
    t = make_trade(WED)
    xl.save_trade(t)
    wb = load_workbook(xl.path)
    ws = wb["Wednesday"]
    row = [c.value for c in ws[2]]
    assert row[:7] == ["T1", "EMA-01", "RELIANCE", "BUY", 10, 1000.0, 1020.0]
    assert row[7] == t.entry_time and row[8] == t.exit_time
    assert row[9] == 200.0
    # charge columns
    assert row[10:17] == [t.brokerage, t.stt, t.transaction_charges, t.sebi_charges, t.gst,
                          t.stamp_duty, t.total_charges]
    assert row[17] == pytest.approx(189.24) and row[18] == "TARGET"
    assert wb["Thursday"].max_row == 1  # other sheets untouched


def test_net_pnl_equals_gross_minus_charges_in_sheet(xl):
    xl.save_trade(make_trade(WED))
    row = [c.value for c in load_workbook(xl.path)["Wednesday"][2]]
    assert row[17] == pytest.approx(row[9] - row[16])


def test_daily_totals_summary(xl):
    win1 = make_trade(WED, exit_=1020.0, tid="T1")
    win2 = make_trade(WED, exit_=1010.0, tid="T2", symbol="TCS")
    loss = make_trade(WED, exit_=990.0, tid="T3", symbol="INFY")
    for t in (win1, win2, loss):
        xl.save_trade(t)
    ws = load_workbook(xl.path)["Wednesday"]
    assert [ws.cell(r, 1).value for r in (2, 3, 4)] == ["T1", "T2", "T3"]
    assert ws.cell(5, 1).value is None  # blank spacer row
    summary = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(6, 12)}
    assert summary["Total Trades"] == 3
    assert summary["Winning Trades"] == 2 and summary["Losing Trades"] == 1
    assert summary["Gross P&L"] == pytest.approx(200 + 100 - 100)
    assert summary["Total Charges"] == pytest.approx(win1.total_charges + win2.total_charges + loss.total_charges)
    assert summary["Net P&L"] == pytest.approx(win1.net_pnl + win2.net_pnl + loss.net_pnl)


def test_trades_on_different_days_go_to_different_sheets(xl):
    xl.save_trade(make_trade(WED, tid="A"))
    xl.save_trade(make_trade(THU, tid="B"))
    wb = load_workbook(xl.path)
    assert wb["Wednesday"].cell(2, 1).value == "A"
    assert wb["Thursday"].cell(2, 1).value == "B"


def test_weekend_trade_is_refused(xl):
    with pytest.raises(ValueError):
        xl.save_trade(make_trade(datetime(2026, 10, 3, 11, 0)))  # Saturday


def test_new_week_archives_previous_workbook(xl, tmp_path):
    xl.save_trade(make_trade(WED, tid="OLD"))
    next_week = datetime(2026, 10, 5, 10, 0)  # following Monday
    xl.save_trade(make_trade(next_week, tid="NEW"))
    archived = tmp_path / "archive" / "weekly_trades_2026-09-28.xlsx"
    assert archived.exists()
    assert load_workbook(archived)["Wednesday"].cell(2, 1).value == "OLD"
    fresh = load_workbook(xl.path)
    assert fresh["Wednesday"].max_row == 1                 # old trade is not in the new week
    assert fresh["Monday"].cell(2, 1).value == "NEW"


def test_archive_never_overwrites_existing_history(xl, tmp_path):
    xl.save_trade(make_trade(WED, tid="A"))
    xl.save_trade(make_trade(datetime(2026, 10, 5, 10, 0), tid="B"))
    # force another rollover that would collide with the same archive name
    wb = load_workbook(xl.path)
    wb.properties.subject = "WEEK_START=2026-09-28"
    wb.save(xl.path)
    xl.save_trade(make_trade(datetime(2026, 10, 6, 10, 0), tid="C"))
    names = sorted(p.name for p in (tmp_path / "archive").iterdir())
    assert names == ["weekly_trades_2026-09-28.xlsx", "weekly_trades_2026-09-28_1.xlsx"]


def test_count_trades(xl):
    assert xl.count_trades(WED.date()) == 0
    xl.save_trade(make_trade(WED, tid="A"))
    xl.save_trade(make_trade(WED, tid="B"))
    assert xl.count_trades(WED.date()) == 2
    assert xl.count_trades(THU.date()) == 0
    assert xl.count_trades(date(2026, 10, 6)) == 0  # different week
