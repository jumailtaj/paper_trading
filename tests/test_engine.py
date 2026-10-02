"""Engine tests driven by simulated ticks (no Kite connection)."""
from datetime import datetime, time, timedelta

import pytest
from openpyxl import load_workbook

from engine.intraday_engine import RUNNING, STOPPED, STOPPING, IntradayEngine
from reporting.weekly_excel import WeeklyExcel
from strategies import registry
from strategies.base import BUY, NO_SIGNAL, SELL, BaseStrategy, Signal
from trading.account import Account
from trading.charge_calculator import ChargeCalculator
from trading.paper_broker import PaperBroker
from trading.positions import PositionManager
from trading.risk import RiskManager

DAY = datetime(2026, 9, 30)  # Wednesday


def at(h, m, s=0):
    return DAY + timedelta(hours=h, minutes=m, seconds=s)


class FakeStrategy:
    """Signals BUY/SELL on the first candle close of each symbol; exit reason is configurable."""
    strategy_id = "FAKE-01"
    timeframe = "5m"
    forced_exit_time = time(15, 15)

    def __init__(self, side=BUY, sl=1.0, rr=2.0):
        self.side, self.sl, self.rr = side, sl, rr
        self.exit_reason = None
        self.calls = 0

    def generate_signal(self, data):
        self.calls += 1
        p = data.ltp
        if self.side == BUY:
            return Signal(self.strategy_id, data.symbol, BUY, p, p - self.sl, p + self.rr * self.sl,
                          timestamp=data.timestamp)
        return Signal(self.strategy_id, data.symbol, SELL, p, p + self.sl, p - self.rr * self.sl,
                      timestamp=data.timestamp)

    def should_exit(self, data, side):
        return self.exit_reason


def build(strategy=None, max_trades=5, excel=None, universe=("AAA", "BBB")):
    strategy = strategy or FakeStrategy()
    return IntradayEngine(strategy, RiskManager(max_trades), PaperBroker(), PositionManager(),
                          Account(100000), ChargeCalculator(), excel, universe=list(universe),
                          position_size=10)


def open_trade(e, sym="AAA", price=100.0):
    """Two ticks across a 5-minute boundary: the second closes the 09:15 candle and triggers the signal."""
    e.handle_tick(sym, price, at(9, 15, 5))
    e.handle_tick(sym, price, at(9, 20, 1))


# ---------------------------------------------------------------- entry
def test_signal_becomes_paper_order_and_position():
    e = build()
    e.start()
    open_trade(e)
    pos = e.positions_view()
    assert len(pos) == 1 and pos[0].symbol == "AAA" and pos[0].entry_price == 100.0
    assert (pos[0].stop_loss, pos[0].target) == (99.0, 102.0)
    orders = e.orders_view()
    assert len(orders) == 1 and orders[0].status == "FILLED" and orders[0].side == "BUY"
    acc = e.account_view()
    assert acc["used_capital"] == 1000 and acc["trades_today"] == 1 and acc["open_positions"] == 1


def test_nothing_happens_when_not_running():
    e = build()
    open_trade(e)
    assert e.positions_view() == [] and e.orders_view() == []


def test_unknown_symbol_and_bad_ticks_ignored():
    e = build()
    e.start()
    e.handle_tick("ZZZ", 100, at(9, 15))
    e.handle_tick("AAA", 0, at(9, 15))
    e.handle_tick("AAA", None, at(9, 15))
    assert e.prices_view() == {"AAA": None, "BBB": None}


# ---------------------------------------------------------------- exits
def test_target_exit_with_charges_and_excel(tmp_path):
    xl = WeeklyExcel(tmp_path / "w.xlsx", tmp_path / "arch")
    e = build(excel=xl)
    e.start()
    open_trade(e)
    e.handle_tick("AAA", 101.0, at(9, 21))
    assert len(e.positions_view()) == 1                      # still open
    e.handle_tick("AAA", 102.0, at(9, 22))
    assert e.positions_view() == []
    t = e.trades_view()[0]
    assert t.exit_reason == "TARGET" and t.gross_pnl == 20.0
    assert t.total_charges > 0 and t.net_pnl == pytest.approx(t.gross_pnl - t.total_charges)
    acc = e.account_view()
    assert acc["equity"] == pytest.approx(100000 + t.net_pnl)
    assert acc["realized_pnl"] == pytest.approx(t.net_pnl) and acc["used_capital"] == 0
    row = [c.value for c in load_workbook(xl.path)["Wednesday"][2]]
    assert row[2] == "AAA" and row[17] == pytest.approx(t.net_pnl) and row[18] == "TARGET"
    assert [o.side for o in e.orders_view()] == ["BUY", "SELL"]


def test_stop_loss_exit():
    e = build()
    e.start()
    open_trade(e)
    e.handle_tick("AAA", 99.0, at(9, 21))
    t = e.trades_view()[0]
    assert t.exit_reason == "STOP_LOSS" and t.gross_pnl == -10.0 and t.net_pnl < -10.0
    assert e.positions_view() == []


def test_strategy_exit_on_candle_close():
    s = FakeStrategy()
    e = build(s)
    e.start()
    open_trade(e)
    s.exit_reason = "OPPOSITE_CROSSOVER"
    e.handle_tick("AAA", 100.5, at(9, 22))                    # same candle: no strategy check yet
    assert len(e.positions_view()) == 1
    e.handle_tick("AAA", 100.5, at(9, 25, 1))                 # next candle closes -> strategy exit
    t = e.trades_view()[0]
    assert t.exit_reason == "OPPOSITE_CROSSOVER" and t.exit_price == 100.5 and t.gross_pnl == 5.0


def test_forced_intraday_exit():
    e = build()
    e.start()
    open_trade(e)
    e.handle_tick("AAA", 100.4, at(15, 14, 59))
    assert len(e.positions_view()) == 1
    e.handle_tick("BBB", 50.0, at(15, 15, 0))                 # a tick for ANY symbol triggers it
    t = e.trades_view()[0]
    assert t.exit_reason == "FORCED_EXIT" and t.exit_price == 100.4
    assert e.positions_view() == []


def test_no_new_entries_after_forced_exit_time():
    s = FakeStrategy()
    e = build(s)
    e.start()
    e.handle_tick("AAA", 100, at(15, 10))
    e.handle_tick("AAA", 100, at(15, 15, 1))
    assert e.positions_view() == [] and s.calls == 0


def test_short_trade_target():
    e = build(FakeStrategy(side=SELL))
    e.start()
    open_trade(e)
    assert e.positions_view()[0].side == "SELL"
    e.handle_tick("AAA", 98.0, at(9, 21))
    t = e.trades_view()[0]
    assert t.exit_reason == "TARGET" and t.gross_pnl == 20.0
    assert [o.side for o in e.orders_view()] == ["SELL", "BUY"]


# ---------------------------------------------------------------- risk integration
def test_max_trades_rejects_signal_and_records_rejected_order():
    e = build(max_trades=1)
    e.start()
    open_trade(e, "AAA")
    open_trade(e, "BBB")
    assert [p.symbol for p in e.positions_view()] == ["AAA"]
    rejected = [o for o in e.orders_view() if o.status == "REJECTED"]
    assert len(rejected) == 1 and rejected[0].symbol == "BBB"
    assert rejected[0].reason == "Maximum trades per day reached"
    assert any("SIGNAL REJECTED" in ev for ev in e.events_view(20))


def test_daily_counter_resets_next_day():
    e = build(max_trades=1)
    e.start()
    open_trade(e, "AAA")
    e.handle_tick("AAA", 102.0, at(9, 22))                    # close it
    nxt = timedelta(days=1)
    e.handle_tick("BBB", 50, at(9, 15, 5) + nxt)
    e.handle_tick("BBB", 50, at(9, 20, 1) + nxt)
    assert [p.symbol for p in e.positions_view()] == ["BBB"]  # allowed again on a new day


# ---------------------------------------------------------------- stop trading
def test_stop_with_no_positions_stops_immediately_and_calls_back():
    e = build()
    called = []
    e.on_stopped = lambda: called.append(1)
    e.start()
    assert e.status == RUNNING
    assert e.stop() == STOPPED and called == [1]


def test_stop_blocks_new_entries_but_keeps_managing_positions():
    e = build()
    called = []
    e.on_stopped = lambda: called.append(1)
    e.start()
    open_trade(e, "AAA")
    assert e.stop() == STOPPING and called == []
    open_trade(e, "BBB")                                      # new signal must not become a trade
    assert [p.symbol for p in e.positions_view()] == ["AAA"]
    e.handle_tick("AAA", 99.0, at(9, 30))                     # existing position is still managed
    assert e.trades_view()[0].exit_reason == "STOP_LOSS"
    assert e.status == STOPPED and called == [1]


def test_cannot_change_strategy_while_running_or_with_positions():
    e = build()
    e.start()
    with pytest.raises(RuntimeError):
        e.set_strategy(registry.create_strategy("ORB-01"))
    e.stop()
    e.set_strategy(registry.create_strategy("ORB-01"))
    assert e.strategy.strategy_id == "ORB-01"


def test_excel_failure_does_not_crash_engine():
    class Broken:
        def save_trade(self, t):
            raise PermissionError("file is open in Excel")
    e = build(excel=Broken())
    e.start()
    open_trade(e)
    e.handle_tick("AAA", 102.0, at(9, 22))
    assert len(e.trades_view()) == 1
    assert any("EXCEL UPDATE FAILED" in ev for ev in e.events_view(20))


def test_strategy_exception_fails_safe():
    class Boom(FakeStrategy):
        def generate_signal(self, data):
            raise RuntimeError("bug")
    e = build(Boom())
    e.start()
    open_trade(e)
    assert e.positions_view() == [] and e.orders_view() == []


# ---------------------------------------------------------------- full pipeline with a REAL strategy
def test_full_pipeline_ema01_on_simulated_ticks(tmp_path):
    closes = [100 - 0.5 * i for i in range(30)] + [85 + 1.0 * i for i in range(30)]
    xl = WeeklyExcel(tmp_path / "w.xlsx", tmp_path / "arch")
    strategy = registry.create_strategy("EMA-01")
    e = IntradayEngine(strategy, RiskManager(5), PaperBroker(), PositionManager(), Account(100000),
                       ChargeCalculator(), xl, universe=["RELIANCE"], position_size=10)
    e.start()
    for i, c in enumerate(closes):
        start = at(9, 15) + timedelta(minutes=5 * i)
        e.handle_tick("RELIANCE", c, start + timedelta(seconds=2), cum_volume=1000 * (i + 1))
        e.handle_tick("RELIANCE", c, start + timedelta(minutes=4, seconds=58), cum_volume=1000 * (i + 1) + 500)
    trades = e.trades_view()
    assert trades, "EMA-01 should have entered and exited at least once"
    first = trades[0]
    assert first.symbol == "RELIANCE" and first.strategy_id == "EMA-01" and first.side == "BUY"
    assert first.exit_reason in ("TARGET", "STOP_LOSS", "OPPOSITE_CROSSOVER", "FORCED_EXIT")
    assert first.net_pnl == pytest.approx(first.gross_pnl - first.total_charges)
    ws = load_workbook(xl.path)["Wednesday"]
    assert ws.cell(2, 1).value == first.trade_id
    assert all(o.status in ("FILLED", "REJECTED") for o in e.orders_view())
